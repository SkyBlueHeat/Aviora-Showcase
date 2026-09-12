import requests
import openpyxl
import json
from datetime import datetime

# Notion Sync Script
# Syncs Excel data with Notion database using Notion API

class NotionSync:
    def __init__(self, api_key, database_id):
        self.api_key = api_key
        self.database_id = database_id
        self.base_url = "https://api.notion.com/v1"
        self.headers = {
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "Notion-Version": "2022-06-28"
        }
    
    def create_page(self, properties):
        """Create a new page in Notion database"""
        url = f"{self.base_url}/pages"
        data = {
            "parent": {"database_id": self.database_id},
            "properties": properties
        }
        
        response = requests.post(url, headers=self.headers, json=data)
        return response.json()
    
    def update_page(self, page_id, properties):
        """Update an existing page in Notion"""
        url = f"{self.base_url}/pages/{page_id}"
        data = {"properties": properties}
        
        response = requests.patch(url, headers=self.headers, json=data)
        return response.json()
    
    def sync_application(self, company, role, status, date_applied, salary):
        """Sync a job application to Notion"""
        properties = {
            "Company": {
                "title": [{"text": {"content": company}}]
            },
            "Role": {
                "rich_text": [{"text": {"content": role}}]
            },
            "Status": {
                "select": {"name": status}
            },
            "Date Applied": {
                "date": {"start": date_applied}
            },
            "Salary": {
                "number": salary
            }
        }
        
        return self.create_page(properties)
    
    def sync_all_applications(self, file_path):
        """Sync all applications from Excel to Notion"""
        wb = openpyxl.load_workbook(file_path)
        ws = wb['Applications']
        
        synced_count = 0
        for row in range(2, ws.max_row + 1):
            # v3.0: Company=col2, Role=col3, Status=col10, DateApplied=col9, Salary=col7
            company = ws.cell(row=row, column=2).value
            role = ws.cell(row=row, column=3).value
            status = ws.cell(row=row, column=10).value
            date_applied = ws.cell(row=row, column=9).value
            salary = ws.cell(row=row, column=7).value
            
            if company and role:
                date_str = date_applied.strftime('%Y-%m-%d') if date_applied else None
                
                try:
                    result = self.sync_application(
                        company, role, status, date_str, salary
                    )
                    if result.get('id'):
                        synced_count += 1
                        print(f"Synced: {company} - {role}")
                except Exception as e:
                    print(f"Error syncing {company}: {e}")
        
        print(f"\nSynced {synced_count} applications to Notion")
        return synced_count

def main():
    # Configuration
    api_key = "your_notion_integration_token"
    database_id = "your_notion_database_id"
    file_path = 'c:/Users/erkay/Desktop/önemli/Developer_Job_Application_Tracker_PRO.xlsx'
    
    # Initialize sync
    sync = NotionSync(api_key, database_id)
    
    # Sync all applications
    sync.sync_all_applications(file_path)

if __name__ == "__main__":
    main()
