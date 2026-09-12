"""Generate a test fixture workbook for backend_bridge tests.
This creates a separate .xlsx file with sample data — NOT the real workbook.
Do NOT run this against Developer_Job_Application_Tracker_PRO.xlsx.
"""
import openpyxl
import os

FIXTURE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tests", "test_fixture_workbook.xlsx")
FIXTURE_PATH = os.path.normpath(FIXTURE_PATH)


def create_fixture():
    wb = openpyxl.Workbook()

    # --- Applications sheet ---
    ws = wb.active
    ws.title = "Applications"
    app_headers = [
        "Application ID", "Company", "Role", "Location", "Country", "Remote",
        "Salary", "Source", "Date Applied", "Status", "Priority", "Fit Score",
        "Follow-up Date", "Recruiter", "LinkedIn Recruiter", "Interview Date",
        "Offer", "Application URL", "Company Career Page", "Resume Version",
        "Cover Letter Version", "Visa Sponsorship", "Employment Type",
        "Work Authorization", "Referral", "Expected Salary", "Current Stage",
        "Days Since Applied", "Days Until Follow-up", "Interview Count",
        "Favorite", "Archive", "Notes",
    ]
    ws.append(app_headers)

    app_data = [
        ["APP-001", "Tech Corp", "Senior Frontend Developer", "San Francisco, CA", "USA", "Yes",
         "$150k - $180k", "LinkedIn", "2026-06-15", "Technical", "High", 85,
         "2026-07-01", "Jane Smith", "", "", "", "https://techcorp.com/jobs/1",
         "https://techcorp.com/careers", "v2", "v1", "Yes", "Full-time", "Yes",
         "Alex", "$170k", "Technical", None, None, 1, "No", "No",
         "Great team, phone screen passed."],
        ["APP-002", "DataFlow", "Backend Engineer", "Remote", "USA", "Yes",
         "$140k - $170k", "Job Board", "2026-06-20", "Applied", "Medium", 72,
         "2026-06-27", "", "", "", "", "https://dataflow.com/jobs/2",
         "https://dataflow.com/careers", "v1", "", "No", "Full-time", "Yes",
         "", "$160k", "Applied", None, None, 0, "No", "No",
         "Applied through job board."],
        ["APP-003", "CloudNet", "Full Stack Engineer", "New York, NY", "USA", "Hybrid",
         "$160k - $200k", "Referral", "2026-06-10", "Offer", "Critical", 95,
         "2026-06-25", "John Doe", "Yes", "2026-06-22", "$190k",
         "https://cloudnet.com/jobs/3", "https://cloudnet.com/careers",
         "v2", "v1", "Yes", "Full-time", "Yes", "Sarah", "$195k",
         "Offer", None, None, 3, "Yes", "No",
         "Got offer! Need to negotiate."],
        ["APP-004", "StartupX", "Frontend Engineer", "Remote", "USA", "Yes",
         "$130k - $160k", "Company Website", "2026-06-28", "Saved", "Low", 65,
         "", "", "", "", "", "https://startupx.com/jobs/4",
         "https://startupx.com/careers", "", "", "No", "Contract", "Yes",
         "", "$150k", "Saved", None, None, 0, "No", "No",
         "Interesting role, need to tailor resume."],
        ["APP-005", "BigTech Inc", "Senior Software Engineer", "Seattle, WA", "USA", "No",
         "$180k - $220k", "LinkedIn", "2026-05-01", "Rejected", "Low", 60,
         "", "", "", "", "", "https://bigtech.com/jobs/5",
         "https://bigtech.com/careers", "v1", "", "Yes", "Full-time", "Yes",
         "", "$200k", "Rejected", None, None, 2, "No", "No",
         "Rejected after technical round."],
        ["APP-006", "GhostCo", "ML Engineer", "San Francisco, CA", "USA", "Remote",
         "$190k - $240k", "Company Website", "2026-04-15", "Ghosted", "Low", 55,
         "", "", "", "", "", "https://ghostco.com/jobs/6",
         "https://ghostco.com/careers", "v1", "", "No", "Full-time", "No",
         "", "$220k", "Ghosted", None, None, 0, "No", "No",
         "No response after 6 weeks."],
    ]
    for row in app_data:
        ws.append(row)

    # --- Salary Comparison sheet ---
    ws2 = wb.create_sheet("Salary Comparison")
    salary_headers = [
        "Company", "Base Salary", "Bonus", "Equity", "Benefits Value",
        "PTO Days", "Remote", "Visa Sponsorship", "Signing Bonus",
        "Relocation Bonus", "Annual Compensation", "Monthly Income",
        "Hourly Rate", "Total Compensation", "Estimated Tax (25%)",
        "Net Salary", "Decision Score", "Weighted Score", "Offer Ranking",
    ]
    ws2.append(salary_headers)
    ws2.append(["Tech Corp", 150000, 20000, 50000, 15000, "20 days", "Yes", "Yes", 10000, 5000,
                None, None, None, None, None, None, None, None, None])
    ws2.append(["CloudNet", 170000, 25000, 60000, 20000, "25 days", "Hybrid", "Yes", 15000, 0,
                None, None, None, None, None, None, None, None, None])

    # --- Interview Tracker sheet ---
    ws3 = wb.create_sheet("Interview Tracker")
    interview_headers = [
        "Company", "Role", "Interview Stage", "Interviewer", "Interview Date",
        "Interview Number", "Duration (min)", "Preparation Status",
        "Confidence Score", "Difficulty", "Result", "Technical Topics",
        "Coding Questions", "Behavioral Questions", "Next Step", "Notes",
        "Lessons Learned",
    ]
    ws3.append(interview_headers)
    ws3.append(["Tech Corp", "Senior Frontend Developer", "Technical", "Jane Smith",
                "2026-06-20", 1, 60, "Completed", 8, "Medium", "Passed",
                "React, TypeScript", "Binary tree traversal", "Tell me about a challenge",
                "Onsite next", "Went well", "Need more system design prep"])
    ws3.append(["CloudNet", "Full Stack Engineer", "Onsite", "Team Panel",
                "2026-06-22", 3, 180, "Completed", 9, "Hard", "Passed",
                "System design, React", "Graph algorithms", "Leadership experience",
                "Offer extended", "Excellent performance", "Strong system design skills"])

    # --- Cover Letter Tracker sheet ---
    ws4 = wb.create_sheet("Cover Letter Tracker")
    cl_headers = [
        "Company", "Role", "Version", "Date Created", "Date Sent",
        "Template Used", "Customization Level", "Response Received",
        "Interview Scheduled", "Success Rate", "A/B Test Group", "Notes",
    ]
    ws4.append(cl_headers)
    ws4.append(["Tech Corp", "Senior Frontend Developer", "v1.0", "2026-06-14", "2026-06-15",
                "Template A", "Heavy Customization", "Yes", "Yes", None, "A", "Great response"])
    ws4.append(["DataFlow", "Backend Engineer", "v1.0", "2026-06-19", "2026-06-20",
                "Template B", "Light Customization", "No", "No", None, "B", "No response yet"])

    # --- Dashboard sheet (layout, minimal) ---
    ws5 = wb.create_sheet("Dashboard")
    ws5["B1"] = "Developer Job Application Tracker PRO"
    ws5["B2"] = "Executive Dashboard"
    ws5["B3"] = "Total Applications"
    ws5["D3"] = "Active Applications"
    ws5["F3"] = "Interviews"
    ws5["H3"] = "Offers"

    wb.save(FIXTURE_PATH)
    print(f"Fixture workbook created: {FIXTURE_PATH}")
    print(f"Size: {os.path.getsize(FIXTURE_PATH)} bytes")


if __name__ == "__main__":
    create_fixture()
