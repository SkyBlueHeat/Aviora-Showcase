"""Read-only Excel reader for JobTracker PRO backend bridge.
Never writes, never saves. Uses openpyxl read_only=True + data_only=True.
"""
import openpyxl
import os
from typing import Any

# Status values that map to React pipeline columns
VALID_STATUSES = {
    "Saved", "Applied", "Phone Screen", "Technical",
    "Onsite", "Offer", "Rejected", "Ghosted",
}


class ExcelReader:
    """Safe read-only wrapper around openpyxl."""

    def __init__(self, workbook_path: str):
        self.workbook_path = workbook_path
        self._wb = None

    def _open(self):
        """Open workbook in read-only mode. Must be closed after use."""
        if self._wb is not None:
            self._wb.close()
        self._wb = openpyxl.load_workbook(
            self.workbook_path,
            read_only=True,
            data_only=True,
        )
        return self._wb

    def _close(self):
        """Close the workbook if open."""
        if self._wb is not None:
            self._wb.close()
            self._wb = None

    def get_workbook_status(self) -> dict:
        """Return basic workbook metadata without exposing the workbook object."""
        if not os.path.exists(self.workbook_path):
            return {
                "exists": False,
                "sheets": [],
                "error": "Workbook file not found",
            }
        try:
            wb = self._open()
            sheets = wb.sheetnames
            self._close()
            return {
                "exists": True,
                "sheets": sheets,
                "sheet_count": len(sheets),
                "error": None,
            }
        except Exception as e:
            self._close()
            return {
                "exists": True,
                "sheets": [],
                "error": str(e),
            }

    def _read_sheet_rows(self, sheet_name: str, header_row: int = 1) -> tuple:
        """Read a sheet and return (headers, rows).
        
        Args:
            sheet_name: Name of the sheet to read.
            header_row: 1-indexed row number that contains headers.
            
        Returns:
            Tuple of (headers list, list of row lists).
            Returns ([], []) if sheet doesn't exist or is empty.
        """
        try:
            wb = self._open()
        except Exception:
            self._close()
            return ([], [])

        if sheet_name not in wb.sheetnames:
            self._close()
            return ([], [])

        ws = wb[sheet_name]
        headers = []
        rows = []
        row_count = 0

        for row in ws.iter_rows(min_row=1, values_only=True):
            row_count += 1
            row_list = list(row)
            if row_count == header_row:
                headers = [str(v).strip() if v is not None else "" for v in row_list]
            elif row_count > header_row:
                is_empty = all(v is None or str(v).strip() == "" for v in row_list)
                if not is_empty:
                    rows.append(row_list)

        self._close()
        return (headers, rows)

    def read_applications(self) -> tuple:
        """Read the Applications sheet. Returns (headers, rows)."""
        return self._read_sheet_rows("Applications", header_row=1)

    def read_salary_comparison(self) -> tuple:
        """Read the Salary Comparison sheet. Returns (headers, rows)."""
        return self._read_sheet_rows("Salary Comparison", header_row=1)

    def read_interview_history(self) -> tuple:
        """Read the Interview Tracker sheet. Returns (headers, rows)."""
        return self._read_sheet_rows("Interview Tracker", header_row=1)

    def read_cover_letter_history(self) -> tuple:
        """Read the Cover Letter Tracker sheet. Returns (headers, rows)."""
        return self._read_sheet_rows("Cover Letter Tracker", header_row=1)

    def read_target_companies(self) -> tuple:
        """Read the Target Companies sheet. Returns (headers, rows)."""
        return self._read_sheet_rows("Target Companies", header_row=1)

    def read_dashboard_summary(self) -> dict:
        """Read the Dashboard sheet (layout sheet — returns label info only)."""
        try:
            wb = self._open()
        except Exception:
            self._close()
            return {"labels": [], "error": "Could not open workbook"}

        if "Dashboard" not in wb.sheetnames:
            self._close()
            return {"labels": [], "error": "Dashboard sheet not found"}

        ws = wb["Dashboard"]
        labels = []
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            for v in row:
                if v is not None and str(v).strip():
                    labels.append(str(v).strip())

        self._close()
        return {"labels": labels, "error": None}
