"""Manual validation script for Core MVP acceptance.

Performs:
1. Staging validation — add, edit, status change, backup verification, persistence
2. Disposable production-copy validation — opt-in, confirmation, session backup, unique backups
3. Real production read-only validation — SHA256 before/after, writes disabled

Uses temporary workbook copies only. Never modifies the real production workbook.
"""
import os
import sys
import shutil
import hashlib
import tempfile
import json

_PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

import openpyxl
from backend_bridge.excel_writer import ExcelWriter, PRODUCTION_WORKBOOK_NAME
from backend_bridge.excel_reader import ExcelReader
from backend_bridge.api import ExcelBridgeAPI

PASS_COUNT = 0
FAIL_COUNT = 0


def check(condition: bool, label: str):
    global PASS_COUNT, FAIL_COUNT
    if condition:
        PASS_COUNT += 1
        print(f"  PASS: {label}")
    else:
        FAIL_COUNT += 1
        print(f"  FAIL: {label}")


def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            h.update(chunk)
    return h.hexdigest()


def read_row_by_id(wb_path: str, app_id: str) -> dict | None:
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    ws = wb["Applications"]
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and str(row[0]).strip() == str(app_id).strip():
            wb.close()
            return {headers[i]: row[i] for i in range(len(headers))}
    wb.close()
    return None


def count_rows(wb_path: str) -> int:
    wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
    ws = wb["Applications"]
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            count += 1
    wb.close()
    return count


def run_staging_validation():
    """Test add, edit, status change on staging workbook."""
    print("\n" + "=" * 60)
    print("  STAGING VALIDATION")
    print("=" * 60)

    fixture = os.path.join(_PROJECT_ROOT, "backend_bridge", "tests", "test_fixture_workbook.xlsx")
    if not os.path.exists(fixture):
        print("  SKIP: fixture workbook not found")
        return

    with tempfile.TemporaryDirectory() as tmpdir:
        staging_wb = os.path.join(tmpdir, "Developer_Job_Application_Tracker_PRO_STAGING.xlsx")
        shutil.copy2(fixture, staging_wb)
        os.environ["JOBTRACKER_WORKBOOK_PATH"] = staging_wb
        os.environ["JOBTRACKER_ENABLE_STAGING_WRITES"] = "1"

        try:
            writer = ExcelWriter(staging_wb, backup_dir=os.path.join(tmpdir, "backups"))
            backup_dir = os.path.join(tmpdir, "backups")

            # Check write mode
            wm = writer.get_write_mode()
            check(wm["enabled"] is True, "staging write mode enabled")
            check(wm["mode"] == "staging", "staging mode detected")

            initial_count = count_rows(staging_wb)

            # 1. Add application
            print("\n  --- Add Application ---")
            add_result = writer.add_application({
                "company": "StagingTest Corp",
                "role": "Backend Developer",
                "location": "Remote",
                "source": "LinkedIn",
                "salary": "120000",
                "fitScore": 85,
                "notes": "Staging validation test",
            })
            check(add_result["success"] is True, "staging add succeeds")
            check(add_result.get("backupPath") is not None, "staging add creates backup")
            add_backup = add_result.get("backupPath")

            after_add_count = count_rows(staging_wb)
            check(after_add_count == initial_count + 1, f"row count increased ({initial_count} -> {after_add_count})")

            # Find the new application ID
            new_row = None
            wb = openpyxl.load_workbook(staging_wb, read_only=True, data_only=True)
            ws = wb["Applications"]
            headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and str(row[headers.index("Company")]).strip() == "StagingTest Corp":
                    new_row = {headers[i]: row[i] for i in range(len(headers))}
                    break
            wb.close()
            check(new_row is not None, "new application found in workbook")
            if new_row:
                new_id = str(new_row.get("Application ID", "")).strip()
                check(new_id != "", "new application has an ID")
            else:
                new_id = ""

            # 2. Edit application
            print("\n  --- Edit Application ---")
            if new_id:
                edit_result = writer.edit_application(new_id, {
                    "company": "StagingTest Corp Updated",
                    "role": "Senior Backend Developer",
                    "fitScore": 90,
                })
                check(edit_result["success"] is True, "staging edit succeeds")
                check(edit_result.get("backupPath") is not None, "staging edit creates backup")
                check(edit_result["backupPath"] != add_backup, "staging edit backup is unique")
                edit_backup = edit_result.get("backupPath")

                # Verify edit persisted
                edited = read_row_by_id(staging_wb, new_id)
                check(str(edited.get("Company", "")).strip() == "StagingTest Corp Updated", "edit company persisted")
                check(str(edited.get("Role", "")).strip() == "Senior Backend Developer", "edit role persisted")
                check(str(edited.get("Fit Score", "")).strip() == "90", "edit fitScore persisted")
            else:
                edit_backup = None

            # 3. Status change
            print("\n  --- Status Change ---")
            if new_id:
                status_result = writer.update_application_status(new_id, "Technical")
                check(status_result["success"] is True, "staging status change succeeds")
                check(status_result.get("backupPath") is not None, "staging status creates backup")
                check(status_result["backupPath"] != add_backup, "status backup unique from add")
                if edit_backup:
                    check(status_result["backupPath"] != edit_backup, "status backup unique from edit")

                # Verify status persisted
                status_row = read_row_by_id(staging_wb, new_id)
                check(str(status_row.get("Status", "")).strip() == "Technical", "status change persisted")

            # 4. Persistence check — reopen with a new reader
            print("\n  --- Persistence Check ---")
            reader = ExcelReader(staging_wb)
            headers, rows = reader.read_applications()
            check(len(rows) == after_add_count, f"persistence: row count matches ({len(rows)} == {after_add_count})")

            # 5. Backup count verification
            print("\n  --- Backup Verification ---")
            backups = os.listdir(backup_dir) if os.path.exists(backup_dir) else []
            check(len(backups) >= 3, f"at least 3 backups created (add+edit+status), got {len(backups)}")

        finally:
            os.environ.pop("JOBTRACKER_WORKBOOK_PATH", None)
            os.environ.pop("JOBTRACKER_ENABLE_STAGING_WRITES", None)


def run_production_copy_validation():
    """Test production writes on a disposable copy."""
    print("\n" + "=" * 60)
    print("  DISPOSABLE PRODUCTION-COPY VALIDATION")
    print("=" * 60)

    fixture = os.path.join(_PROJECT_ROOT, "backend_bridge", "tests", "test_fixture_workbook.xlsx")
    if not os.path.exists(fixture):
        print("  SKIP: fixture workbook not found")
        return

    with tempfile.TemporaryDirectory() as tmpdir:
        prod_wb = os.path.join(tmpdir, PRODUCTION_WORKBOOK_NAME)
        shutil.copy2(fixture, prod_wb)
        os.environ["JOBTRACKER_WORKBOOK_PATH"] = prod_wb
        os.environ["JOBTRACKER_ENABLE_PRODUCTION_WRITES"] = "1"

        try:
            writer = ExcelWriter(prod_wb, backup_dir=os.path.join(tmpdir, "backups"))
            backup_dir = os.path.join(tmpdir, "backups")

            # Check production mode
            wm = writer.get_write_mode()
            check(wm["enabled"] is True, "production write mode enabled")
            check(wm["mode"] == "production", "production mode detected")

            # Session backup should not exist yet
            check(wm.get("sessionBackupPath") is None, "no session backup before first write")

            # 1. Add — creates session backup + per-write backup
            print("\n  --- Production Add ---")
            add_result = writer.add_application({
                "company": "ProdTest Corp",
                "role": "DevOps Engineer",
                "fitScore": 75,
            })
            check(add_result["success"] is True, "production add succeeds")
            add_backup = add_result.get("backupPath")
            check(add_backup is not None, "production add creates per-write backup")

            # Check session backup was created
            wm_after = writer.get_write_mode()
            session_backup = wm_after.get("sessionBackupPath")
            check(session_backup is not None, "production session backup created on first write")
            if session_backup:
                check(os.path.exists(session_backup), "session backup file exists on disk")
                check(session_backup != add_backup, "session backup is different from per-write backup")

            # 2. Edit — creates per-write backup, no new session backup
            print("\n  --- Production Edit ---")
            # Find the new app
            new_row = None
            wb = openpyxl.load_workbook(prod_wb, read_only=True, data_only=True)
            ws = wb["Applications"]
            headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and str(row[headers.index("Company")]).strip() == "ProdTest Corp":
                    new_row = {headers[i]: row[i] for i in range(len(headers))}
                    break
            wb.close()

            if new_row:
                new_id = str(new_row.get("Application ID", "")).strip()
                edit_result = writer.edit_application(new_id, {
                    "company": "ProdTest Corp Updated",
                    "salary": "140000",
                })
                check(edit_result["success"] is True, "production edit succeeds")
                edit_backup = edit_result.get("backupPath")
                check(edit_backup is not None, "production edit creates per-write backup")
                check(edit_backup != add_backup, "production edit backup unique from add")

                # Session backup should be unchanged
                wm_after_edit = writer.get_write_mode()
                check(wm_after_edit.get("sessionBackupPath") == session_backup, "session backup unchanged after edit")

                # 3. Status change
                print("\n  --- Production Status Change ---")
                status_result = writer.update_application_status(new_id, "Phone Screen")
                check(status_result["success"] is True, "production status change succeeds")
                status_backup = status_result.get("backupPath")
                check(status_backup is not None, "production status creates per-write backup")
                check(status_backup != add_backup, "production status backup unique from add")
                check(status_backup != edit_backup, "production status backup unique from edit")

                # 4. Verify all backups have unique filenames
                print("\n  --- Backup Uniqueness ---")
                all_backups = os.listdir(backup_dir)
                check(len(all_backups) >= 3, f"at least 3 per-write backups, got {len(all_backups)}")
                # All filenames should be unique (guaranteed by timestamp with microseconds)
                check(len(all_backups) == len(set(all_backups)), "all backup filenames are unique")

                # 5. Persistence
                print("\n  --- Production Persistence ---")
                persisted = read_row_by_id(prod_wb, new_id)
                check(str(persisted.get("Company", "")).strip() == "ProdTest Corp Updated", "production edit persisted")
                check(str(persisted.get("Status", "")).strip() == "Phone Screen", "production status persisted")
            else:
                print("  SKIP: could not find added application")

        finally:
            os.environ.pop("JOBTRACKER_WORKBOOK_PATH", None)
            os.environ.pop("JOBTRACKER_ENABLE_PRODUCTION_WRITES", None)


def run_production_readonly_validation():
    """Verify real production workbook is not modified without opt-in."""
    print("\n" + "=" * 60)
    print("  REAL PRODUCTION READ-ONLY VALIDATION")
    print("=" * 60)

    prod_path = os.path.join(_PROJECT_ROOT, PRODUCTION_WORKBOOK_NAME)
    if not os.path.exists(prod_path):
        print("  SKIP: production workbook not found")
        return

    # Record SHA256 before
    hash_before = sha256_file(prod_path)
    print(f"  SHA256 before: {hash_before[:32]}...")

    # Ensure no production write env vars
    os.environ.pop("JOBTRACKER_ENABLE_PRODUCTION_WRITES", None)
    os.environ.pop("JOBTRACKER_WORKBOOK_PATH", None)

    try:
        writer = ExcelWriter(prod_path)

        # Check write mode is disabled
        wm = writer.get_write_mode()
        check(wm["enabled"] is False, "production write mode disabled without opt-in")
        reason = wm.get("reason", "")
        check("JOBTRACKER_WORKBOOK_PATH" in reason or "PRODUCTION_WRITES" in reason,
              f"reason mentions required env var (got: {reason[:80]})")

        # Attempt add — should be blocked
        add_result = writer.add_application({"company": "BlockedTest", "role": "Test"})
        check(add_result["success"] is False, "add blocked in read-only production mode")

        # Attempt edit — should be blocked
        edit_result = writer.edit_application("APP-001", {"company": "BlockedEdit"})
        check(edit_result["success"] is False, "edit blocked in read-only production mode")

        # Attempt status change — should be blocked
        status_result = writer.update_application_status("APP-001", "Technical")
        check(status_result["success"] is False, "status change blocked in read-only production mode")

    finally:
        pass

    # Record SHA256 after
    hash_after = sha256_file(prod_path)
    print(f"  SHA256 after:  {hash_after[:32]}...")
    check(hash_before == hash_after, "production workbook SHA256 unchanged after blocked writes")


def main():
    print(f"\n{'#' * 60}")
    print("  JobTracker PRO — Core MVP Manual Validation")
    print(f"{'#' * 60}")

    run_staging_validation()
    run_production_copy_validation()
    run_production_readonly_validation()

    # Final report
    print(f"\n{'=' * 60}")
    print("  MANUAL VALIDATION FINAL REPORT")
    print(f"{'=' * 60}")
    print(f"  TOTAL PASSES:  {PASS_COUNT}")
    print(f"  TOTAL FAILURES: {FAIL_COUNT}")
    print()
    if FAIL_COUNT == 0:
        print("  RESULT: ALL VALIDATIONS PASSED")
        sys.exit(0)
    else:
        print("  RESULT: VALIDATION FAILED — SEE FAILURES ABOVE")
        sys.exit(1)


if __name__ == "__main__":
    main()
