import os
import openpyxl
import csv
import json
import html as html_lib
from datetime import datetime
import pandas as pd

# Multi-format Data Export Script
# Exports Excel data to CSV, JSON, HTML, and PDF formats

def export_to_csv(ws, output_path):
    """Export worksheet to CSV"""
    with open(output_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(row)
    print(f"CSV export completed: {output_path}")

def export_to_json(ws, output_path):
    """Export worksheet to JSON"""
    data = []
    headers = [cell.value for cell in ws[1]]
    
    for row in ws.iter_rows(min_row=2):
        row_data = {}
        for i, cell in enumerate(row):
            if i < len(headers):
                value = cell.value
                if isinstance(value, datetime):
                    value = value.strftime('%Y-%m-%d')
                row_data[headers[i]] = value
        data.append(row_data)
    
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, indent=2, default=str)
    print(f"JSON export completed: {output_path}")

def export_to_html(ws, output_path):
    """Export worksheet to HTML"""
    html = """
    <!DOCTYPE html>
    <html>
    <head>
        <title>Job Application Data</title>
        <style>
            table { border-collapse: collapse; width: 100%; }
            th, td { border: 1px solid #ddd; padding: 8px; text-align: left; }
            th { background-color: #222222; color: white; }
            tr:nth-child(even) { background-color: #f2f2f2; }
        </style>
    </head>
    <body>
        <h1>Job Application Data</h1>
        <table>
    """
    
    # Headers
    html += "<tr>"
    for cell in ws[1]:
        header_value = "" if cell.value is None else html_lib.escape(str(cell.value))
        html += f"<th>{header_value}</th>"
    html += "</tr>"
    
    # Data rows
    for row in ws.iter_rows(min_row=2):
        html += "<tr>"
        for cell in row:
            value = cell.value
            if isinstance(value, datetime):
                value = value.strftime('%Y-%m-%d')
            cell_value = "" if value is None else html_lib.escape(str(value))
            html += f"<td>{cell_value}</td>"
        html += "</tr>"
    
    html += """
        </table>
    </body>
    </html>
    """
    
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f"HTML export completed: {output_path}")

def export_to_pdf(ws, output_path):
    """Export worksheet to PDF using pandas and matplotlib"""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.backends.backend_pdf import PdfPages
        
        # Convert to DataFrame
        data = []
        headers = [cell.value for cell in ws[1]]
        
        for row in ws.iter_rows(min_row=2):
            row_data = []
            for cell in row:
                value = cell.value
                if isinstance(value, datetime):
                    value = value.strftime('%Y-%m-%d')
                row_data.append(value)
            data.append(row_data)
        
        df = pd.DataFrame(data, columns=headers)
        
        # Create PDF
        with PdfPages(output_path) as pdf:
            fig, ax = plt.subplots(figsize=(12, 8))
            ax.axis('tight')
            ax.axis('off')
            
            table = ax.table(cellText=df.values, colLabels=df.columns,
                           cellLoc='center', loc='center')
            table.auto_set_font_size(False)
            table.set_fontsize(8)
            table.scale(1, 2)
            
            pdf.savefig(fig, bbox_inches='tight')
            plt.close()
        
        print(f"PDF export completed: {output_path}")
    except ImportError:
        print("PDF export requires matplotlib. Install with: pip install matplotlib")

def export_all_sheets(file_path, output_dir):
    """Export all sheets to multiple formats"""
    wb = openpyxl.load_workbook(file_path)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        safe_name = sheet_name.replace(' ', '_').replace('/', '_')
        
        # Create output directory
        os.makedirs(output_dir, exist_ok=True)
        
        # Export to CSV
        csv_path = f"{output_dir}/{safe_name}_{timestamp}.csv"
        export_to_csv(ws, csv_path)
        
        # Export to JSON
        json_path = f"{output_dir}/{safe_name}_{timestamp}.json"
        export_to_json(ws, json_path)
        
        # Export to HTML
        html_path = f"{output_dir}/{safe_name}_{timestamp}.html"
        export_to_html(ws, html_path)
        
        # Export to PDF
        pdf_path = f"{output_dir}/{safe_name}_{timestamp}.pdf"
        export_to_pdf(ws, pdf_path)

def main():
    _base = os.path.dirname(os.path.abspath(__file__))
    file_path = os.path.join(_base, 'Developer_Job_Application_Tracker_PRO.xlsx')
    output_dir = os.path.join(_base, 'exports')
    
    export_all_sheets(file_path, output_dir)
    print("\nAll exports completed successfully!")

if __name__ == "__main__":
    import os
    main()
