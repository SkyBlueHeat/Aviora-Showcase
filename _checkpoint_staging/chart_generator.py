import openpyxl
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
import matplotlib.pyplot as plt
import pandas as pd

# Chart Generator Script
# Creates visual charts from job application data

def create_pipeline_chart(file_path, output_path):
    """Create a pipeline status pie chart"""
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Applications']
    
    # Count statuses
    status_counts = {}
    for row in range(2, ws.max_row + 1):
        # v3.0: Status=col10
        status = ws.cell(row=row, column=10).value
        if status:
            status_counts[status] = status_counts.get(status, 0) + 1
    
    # Create matplotlib pie chart
    if status_counts:
        labels = list(status_counts.keys())
        sizes = list(status_counts.values())
        colors = ['#FF6B35', '#20C4B7', '#FFD700', '#FF6347', '#9370DB', '#32CD32']
        
        plt.figure(figsize=(10, 8))
        plt.pie(sizes, labels=labels, colors=colors[:len(labels)], autopct='%1.1f%%',
                startangle=90, textprops={'fontsize': 12})
        plt.title('Application Pipeline Status', fontsize=16, fontweight='bold')
        plt.axis('equal')
        plt.savefig(output_path, dpi=300, bbox_inches='tight')
        plt.close()
        print(f"Pipeline chart created: {output_path}")

def create_weekly_trend_chart(file_path, output_path):
    """Create weekly application trend line chart"""
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Applications']
    
    # Count applications by week
    weekly_counts = {}
    for row in range(2, ws.max_row + 1):
        # v3.0: DateApplied=col9
        date = ws.cell(row=row, column=9).value
        if date and isinstance(date, type(openpyxl.utils.datetime.datetime.now())):
            week = date.strftime('%Y-W%W')
            weekly_counts[week] = weekly_counts.get(week, 0) + 1
    
    if weekly_counts:
        weeks = sorted(weekly_counts.keys())
        counts = [weekly_counts[week] for week in weeks]
        
        plt.figure(figsize=(12, 6))
        plt.plot(weeks, counts, marker='o', linewidth=2, markersize=8, color='#FF6B35')
        plt.title('Weekly Application Trend', fontsize=16, fontweight='bold')
        plt.xlabel('Week', fontsize=12)
        plt.ylabel('Applications', fontsize=12)
        plt.xticks(rotation=45, ha='right')
        plt.grid(True, alpha=0.3)
        plt.tight_layout()
        plt.savefig(output_path, dpi=300, bbox_inches='tight')
        plt.close()
        print(f"Weekly trend chart created: {output_path}")

def create_salary_distribution_chart(file_path, output_path):
    """Create salary distribution bar chart"""
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Applications']
    
    # Collect salaries
    salaries = []
    for row in range(2, ws.max_row + 1):
        # v3.0: Salary=col7
        salary = ws.cell(row=row, column=7).value
        if salary and isinstance(salary, (int, float)):
            salaries.append(salary)
    
    if salaries:
        # Create salary ranges
        ranges = ['0-50k', '50-100k', '100-150k', '150-200k', '200k+']
        range_counts = [0, 0, 0, 0, 0]
        
        for salary in salaries:
            if salary < 50000:
                range_counts[0] += 1
            elif salary < 100000:
                range_counts[1] += 1
            elif salary < 150000:
                range_counts[2] += 1
            elif salary < 200000:
                range_counts[3] += 1
            else:
                range_counts[4] += 1
        
        plt.figure(figsize=(10, 6))
        bars = plt.bar(ranges, range_counts, color='#20C4B7', edgecolor='#222222', linewidth=2)
        plt.title('Salary Distribution', fontsize=16, fontweight='bold')
        plt.xlabel('Salary Range', fontsize=12)
        plt.ylabel('Count', fontsize=12)
        plt.grid(True, alpha=0.3, axis='y')
        
        # Add value labels on bars
        for bar in bars:
            height = bar.get_height()
            plt.text(bar.get_x() + bar.get_width()/2., height,
                    f'{int(height)}',
                    ha='center', va='bottom', fontsize=11)
        
        plt.tight_layout()
        plt.savefig(output_path, dpi=300, bbox_inches='tight')
        plt.close()
        print(f"Salary distribution chart created: {output_path}")

def create_response_rate_chart(file_path, output_path):
    """Create response rate comparison chart"""
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Applications']
    
    total = 0
    responses = 0
    
    for row in range(2, ws.max_row + 1):
        # v3.0: Status=col10
        status = ws.cell(row=row, column=10).value
        if status:
            total += 1
            if status in ['Interview', 'Offer', 'Rejected']:
                responses += 1
    
    if total > 0:
        response_rate = (responses / total) * 100
        no_response = 100 - response_rate
        
        labels = ['Response Received', 'No Response']
        sizes = [response_rate, no_response]
        colors = ['#20C4B7', '#FF6B35']
        
        plt.figure(figsize=(8, 8))
        plt.pie(sizes, labels=labels, colors=colors, autopct='%1.1f%%',
                startangle=90, textprops={'fontsize': 14})
        plt.title(f'Response Rate ({response_rate:.1f}%)', fontsize=16, fontweight='bold')
        plt.axis('equal')
        plt.savefig(output_path, dpi=300, bbox_inches='tight')
        plt.close()
        print(f"Response rate chart created: {output_path}")

def create_all_charts(file_path, output_dir):
    """Generate all charts"""
    import os
    os.makedirs(output_dir, exist_ok=True)
    
    timestamp = pd.Timestamp.now().strftime("%Y%m%d_%H%M%S")
    
    create_pipeline_chart(file_path, f"{output_dir}/pipeline_{timestamp}.png")
    create_weekly_trend_chart(file_path, f"{output_dir}/weekly_trend_{timestamp}.png")
    create_salary_distribution_chart(file_path, f"{output_dir}/salary_distribution_{timestamp}.png")
    create_response_rate_chart(file_path, f"{output_dir}/response_rate_{timestamp}.png")
    
    print("\nAll charts generated successfully!")

def main():
    import os as _os
    _base = _os.path.dirname(_os.path.abspath(__file__))
    file_path = _os.path.join(_base, 'Developer_Job_Application_Tracker_PRO.xlsx')
    output_dir = _os.path.join(_base, 'charts')
    
    create_all_charts(file_path, output_dir)

if __name__ == "__main__":
    main()
