import openpyxl
from datetime import date, datetime
import re

# Data Validation Script
# Validates data integrity and consistency across all sheets


def _is_valid_date_value(value):
    if value in (None, ""):
        return True
    if isinstance(value, (datetime, date)):
        return True
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return True
        for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%m/%d/%Y', '%d/%m/%Y'):
            try:
                datetime.strptime(text, fmt)
                return True
            except ValueError:
                pass
    return False

def validate_applications_sheet(file_path):
    """Validate Applications sheet data"""
    print("Validating Applications Sheet...")
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Applications']
    
    errors = []
    warnings = []
    
    for row in range(2, ws.max_row + 1):
        # v3.0: Company=col2, Role=col3, Status=col10, DateApplied=col9, Salary=col7
        company = ws.cell(row=row, column=2).value
        role = ws.cell(row=row, column=3).value
        status = ws.cell(row=row, column=10).value
        date_applied = ws.cell(row=row, column=9).value
        salary = ws.cell(row=row, column=7).value
        
        # Check required fields
        if not company and not role:
            continue  # Empty row
        
        if not company:
            errors.append(f"Row {row}: Missing Company")
        
        if not role:
            errors.append(f"Row {row}: Missing Role")
        
        if not status:
            warnings.append(f"Row {row}: Missing Status")
        
        # Validate status values
        valid_statuses = ['Applied', 'Screening', 'Interview', 'Offer', 'Rejected', 'Ghosted', 'Withdrawn', 'Declined', 'Accepted', 'Hired']
        if status and status not in valid_statuses:
            errors.append(f"Row {row}: Invalid status '{status}'")
        
        # Validate date format
        if not _is_valid_date_value(date_applied):
            errors.append(f"Row {row}: Invalid date format")
        
        # Validate salary
        if salary and not isinstance(salary, (int, float)):
            errors.append(f"Row {row}: Invalid salary format")
        elif salary and salary < 0:
            errors.append(f"Row {row}: Negative salary")
    
    print(f"  Errors: {len(errors)}")
    print(f"  Warnings: {len(warnings)}")
    
    for error in errors[:5]:  # Show first 5 errors
        print(f"  ✗ {error}")
    
    return errors, warnings

def validate_interview_tracker(file_path):
    """Validate Interview Tracker data"""
    print("Validating Interview Tracker...")
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Interview Tracker']
    
    errors = []
    
    for row in range(2, ws.max_row + 1):
        # v3.0: Company=col1, Stage=col3, Result=col11
        company = ws.cell(row=row, column=1).value
        stage = ws.cell(row=row, column=3).value
        result = ws.cell(row=row, column=11).value
        
        if not company:
            continue
        
        # Validate stage
        valid_stages = ['Recruiter Screen', 'Technical', 'Live Coding', 'System Design', 'HR', 'Final', 'Offer', 'Rejected']
        if stage and stage not in valid_stages:
            errors.append(f"Row {row}: Invalid interview stage '{stage}'")
        
        # Validate result
        valid_results = ['Pending', 'Passed', 'Rejected']
        if result and result not in valid_results:
            errors.append(f"Row {row}: Invalid result '{result}'")
    
    print(f"  Errors: {len(errors)}")
    
    return errors

def validate_salary_comparison(file_path):
    """Validate Salary Comparison data"""
    print("Validating Salary Comparison...")
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Salary Comparison']
    
    errors = []
    
    for row in range(2, ws.max_row + 1):
        company = ws.cell(row=row, column=1).value
        base_salary = ws.cell(row=row, column=2).value
        bonus = ws.cell(row=row, column=3).value
        equity = ws.cell(row=row, column=4).value
        
        if not company:
            continue
        
        # Validate numeric fields
        if base_salary and not isinstance(base_salary, (int, float)):
            errors.append(f"Row {row}: Invalid base salary")
        
        if bonus and not isinstance(bonus, (int, float)):
            errors.append(f"Row {row}: Invalid bonus")
        
        if equity and not isinstance(equity, (int, float)):
            errors.append(f"Row {row}: Invalid equity")
    
    print(f"  Errors: {len(errors)}")
    
    return errors

def validate_email_format(file_path):
    """Validate email formats in Recruiter CRM"""
    print("Validating Email Formats...")
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Recruiter CRM']
    
    errors = []
    
    for row in range(2, ws.max_row + 1):
        email = ws.cell(row=row, column=3).value
        
        if email:
            # Simple email validation
            pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
            if not re.match(pattern, email):
                errors.append(f"Row {row}: Invalid email format '{email}'")
    
    print(f"  Errors: {len(errors)}")
    
    return errors

def validate_linkedin_urls(file_path):
    """Validate LinkedIn URL formats"""
    print("Validating LinkedIn URLs...")
    wb = openpyxl.load_workbook(file_path)
    ws = wb['Recruiter CRM']
    
    errors = []
    
    for row in range(2, ws.max_row + 1):
        linkedin = ws.cell(row=row, column=4).value
        
        if linkedin:
            if 'linkedin.com' not in linkedin.lower():
                errors.append(f"Row {row}: Invalid LinkedIn URL '{linkedin}'")
    
    print(f"  Errors: {len(errors)}")
    
    return errors

def generate_validation_report(file_path, output_path):
    """Generate comprehensive validation report"""
    print("="*60)
    print("DATA VALIDATION REPORT")
    print("="*60)
    print()
    
    all_errors = []
    all_warnings = []
    
    # Run all validations
    errors, warnings = validate_applications_sheet(file_path)
    all_errors.extend(errors)
    all_warnings.extend(warnings)
    
    errors = validate_interview_tracker(file_path)
    all_errors.extend(errors)
    
    errors = validate_salary_comparison(file_path)
    all_errors.extend(errors)
    
    errors = validate_email_format(file_path)
    all_errors.extend(errors)
    
    errors = validate_linkedin_urls(file_path)
    all_errors.extend(errors)
    
    # Generate report
    report = f"""
================================================================================
                    DATA VALIDATION REPORT
================================================================================
Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
File: {file_path}

================================================================================
SUMMARY
================================================================================
Total Errors: {len(all_errors)}
Total Warnings: {len(all_warnings)}

================================================================================
DETAILED ERRORS
================================================================================
"""
    
    for error in all_errors:
        report += f"✗ {error}\n"
    
    report += f"""
================================================================================
DETAILED WARNINGS
================================================================================
"""
    
    for warning in all_warnings:
        report += f"⚠ {warning}\n"
    
    if len(all_errors) == 0 and len(all_warnings) == 0:
        report += "\n✓ No data validation issues found. All data is clean!\n"
    else:
        report += f"\n⚠ Please review and fix the {len(all_errors)} error(s) and {len(all_warnings)} warning(s) above.\n"
    
    report += """
================================================================================
CreatorDockStudio - Quality Assurance
================================================================================
"""
    
    # Save report
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(report)
    
    print(f"\nValidation report saved: {output_path}")
    print(f"\nTotal Errors: {len(all_errors)}")
    print(f"Total Warnings: {len(all_warnings)}")

def main():
    import os as _os
    _base = _os.path.dirname(_os.path.abspath(__file__))
    file_path = _os.path.join(_base, 'Developer_Job_Application_Tracker_PRO.xlsx')
    output_dir = _os.path.join(_base, 'reports')
    _os.makedirs(output_dir, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = f"{output_dir}/validation_report_{timestamp}.txt"
    
    generate_validation_report(file_path, output_path)

if __name__ == "__main__":
    main()
