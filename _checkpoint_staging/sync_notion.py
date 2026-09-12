import os
from datetime import datetime, date

import openpyxl
import requests


class NotionSync:
    def __init__(self, api_key: str, database_id: str):
        self.api_key = api_key.strip()
        self.database_id = database_id.strip()
        self.base_url = "https://api.notion.com/v1"
        self.headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json",
            "Notion-Version": "2022-06-28",
        }

    def create_page(self, properties: dict) -> dict:
        url = f"{self.base_url}/pages"
        payload = {
            "parent": {"database_id": self.database_id},
            "properties": properties,
        }
        response = requests.post(url, headers=self.headers, json=payload, timeout=30)
        response.raise_for_status()
        return response.json()

    def sync_application(self, company, role, status, date_applied, salary) -> dict:
        properties = {
            "Company": {
                "title": [{"text": {"content": str(company or "")[:2000]}}]
            },
            "Role": {
                "rich_text": [{"text": {"content": str(role or "")[:2000]}}]
            },
            "Status": {
                "select": {"name": str(status or "Applied")[:100]}
            },
        }

        if date_applied:
            properties["Date Applied"] = {"date": {"start": date_applied}}
        if salary not in (None, ""):
            try:
                properties["Salary"] = {"number": float(salary)}
            except Exception:
                pass

        return self.create_page(properties)

    def sync_all_applications(self, file_path: str) -> int:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        if "Applications" not in wb.sheetnames:
            raise ValueError("Applications sheet not found in tracker.")
        ws = wb["Applications"]

        synced_count = 0
        for row in range(2, ws.max_row + 1):
            # v3.0: Company=col2, Role=col3, Salary=col7, DateApplied=col9, Status=col10
            company = ws.cell(row=row, column=2).value
            role = ws.cell(row=row, column=3).value
            salary = ws.cell(row=row, column=7).value
            date_applied = ws.cell(row=row, column=9).value
            status = ws.cell(row=row, column=10).value

            if not company or not role:
                continue

            date_str = _normalize_date(date_applied)

            try:
                result = self.sync_application(company, role, status, date_str, salary)
                if result.get("id"):
                    synced_count += 1
                    print(f"[OK] Synced: {company} - {role}")
                else:
                    print(f"[!] Notion did not return a page id for: {company} - {role}")
            except requests.HTTPError as exc:
                detail = _safe_response_text(exc.response)
                print(f"[!] HTTP error syncing {company}: {exc}\n{detail}")
            except Exception as exc:
                print(f"[!] Error syncing {company}: {exc}")

        print(f"\nSynced {synced_count} applications to Notion")
        return synced_count


def _normalize_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except Exception:
            pass
    return None


def _safe_response_text(response):
    if response is None:
        return ""
    try:
        return response.text
    except Exception:
        return ""


def sync_to_notion(file_path: str, api_key: str, database_id: str) -> int:
    if not api_key or not database_id:
        raise ValueError("Notion API key and database ID are required.")
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Tracker file not found: {file_path}")
    syncer = NotionSync(api_key, database_id)
    return syncer.sync_all_applications(file_path)


def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    file_path = os.path.join(base_dir, "Developer_Job_Application_Tracker_PRO.xlsx")
    api_key = os.environ.get("NOTION_API_KEY", "").strip()
    database_id = os.environ.get("NOTION_DATABASE_ID", "").strip()
    if not api_key or not database_id:
        raise SystemExit("Set NOTION_API_KEY and NOTION_DATABASE_ID environment variables before running.")
    sync_to_notion(file_path, api_key, database_id)


if __name__ == "__main__":
    main()
